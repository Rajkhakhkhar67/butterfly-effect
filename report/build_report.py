from pathlib import Path
from datetime import date
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm, inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, KeepTogether

ROOT = Path(r'D:\butterfly-effect')
OUT = ROOT / 'report'
OUT.mkdir(exist_ok=True)
QA = OUT / 'qa'
QA.mkdir(exist_ok=True)
FONT = r'C:\Windows\Fonts\Nirmala.ttf'

PROMPTS = [
    ('Prompt 1 — diagnosis and repair', Path(r'C:\Users\com\.codex\attachments\332eaed1-9363-4095-92b3-f600c6388d04\Pasted text.txt'), 'The Butterfly Effect project handoff requesting inspection, repair, verification, and a running preview.'),
    ('Prompt 2 — final development pass', Path(r'C:\Users\com\.codex\attachments\c0cb0389-5bdf-4e66-a97c-77223023bc1e\Pasted text.txt'), 'The final application quality, testing, deployment-readiness, and handover request.'),
    ('Prompt 3 — report and prompt book', Path(r'C:\Users\com\.codex\attachments\a3a3a649-fb2b-44f2-a127-a9642cb85a31\Pasted text.txt'), 'The current request to inspect the university format, prepare the report and prompt book, and list student actions.'),
]

PLACEHOLDER = '[STUDENT ACTION REQUIRED: '

report_sections = [
('Abstract', [
    'The Butterfly Effect is a browser-based event-planning scenario demonstrator. It lets a user change six planning inputs and immediately see calculated spending, budget balance, preparation pressure, schedule risk, and a readiness category. A dependency map highlights pathways from changed inputs to modeled outcomes, and a fixed baseline remains available for comparison.',
    'The application uses transparent deterministic rules implemented in TypeScript. It is intended for illustrating a simplified model, not for forecasting real events or proving scientific causality. The project was built with React and Vite and does not require a backend or an API key. Available verification records include a production build, lint run, six simulation tests, and browser interaction checks. The complete university report PDF was not present in the project; the supplied cover image does not expose its full section requirements.']),
('1. Introduction', [
    'Planning an event involves connected decisions. Increasing the guest count changes catering expenditure; expenditure affects the remaining budget; and reducing preparation time raises the model’s pressure and schedule-risk values. A person considering one decision at a time may miss these linked effects.',
    'The Butterfly Effect provides a small interactive environment for exploring those relationships. The application focuses on event planning because costs, guest counts, available budget, and time remaining can be represented as clear inputs and outputs. The model is deliberately compact so its rules can be followed during a demonstration.']),
('1.1 Problem statement', [
    'A static list of event costs does not make dependencies between planning choices and outcomes easy to explore. Users need a way to change assumptions, view recalculated results, and compare the changed scenario with a stable starting point.']),
('1.2 Motivation and intended users', [
    'The project was prepared as an interactive university IPE demonstration. Its intended users are students, instructors, and other viewers who want to see how a small set of scenario assumptions affects a simple planning model. No user study or survey was conducted, so the report makes no claims about measured learning outcomes or user satisfaction.']),
('2. Objectives', [
    'The project objectives are to provide editable event-planning variables; calculate related cost and schedule indicators deterministically; visualize modeled dependencies; compare a current scenario with a fixed baseline; provide quick presets and reset; and make the application usable in a browser without server-side services or API credentials.']),
('3. Proposed solution', [
    'The application presents six controls: venue cost, guest count, catering cost per guest, travel cost, available budget, and days remaining. The simulation computes catering expenditure as guest count multiplied by catering cost per guest. Total expenditure is venue cost plus catering expenditure plus travel cost. Budget remaining equals available budget minus total expenditure; deficit is the positive overspend, or zero when spending stays within the budget.',
    'Preparation pressure is derived from days remaining on a bounded 0–100 scale. Schedule risk is computed from that pressure. Readiness is assigned using the budget balance, schedule-risk value, and days remaining. The interface shows these relationships as an interactive graph and as result cards. Four presets alter one input at a time: double guests, increase venue cost, shorten preparation time, and increase budget.']),
('4. Technology stack', [
    'The project manifest identifies React and React DOM 19.x, TypeScript 7.0.x, Vite 8.3.x, and Oxlint 1.89.x. The application is implemented in TSX and CSS. Node.js’s built-in test runner executes the simulation tests. No additional test framework, backend, paid API, or runtime service is used. Versions above reflect the project manifest and lockfile in the inspected project.']),
('5. System design and architecture', [
    'The interface layer is implemented in src/App.tsx. It owns the current scenario and selected graph node, renders the input controls and presets, calls the simulation functions, and displays results and comparison values. The calculation layer is src/simulation.ts, which contains scenario types, ranges, validation, and deterministic calculations. The test layer is tests/simulation.test.mjs. Vite serves the development application and produces static files in dist for deployment.',
    'The baseline values are a separate fixed object in the simulation module. Editing a current value does not mutate the baseline. Reset copies the baseline into current scenario state. The graph uses an explicit list of modeled edges and highlights downstream nodes from changed or selected nodes.']),
('5.1 Workflow', [
    'The user changes a slider or numeric field. The input is bounded to its allowed range and step. React state updates the current scenario; the deterministic simulation recomputes results; the dashboard and graph update; and the baseline comparison shows the changed input and output deltas. Selecting a graph node highlights its modeled downstream path.']),
('6. Implemented features', [
    'Scenario controls with six labeled range sliders and numeric inputs.',
    'Calculations for catering expenditure, total expenditure, budget remaining, budget deficit, preparation pressure, schedule risk, and readiness.',
    'An interactive SVG cause-and-effect graph with keyboard-operable node controls and downstream highlighting.',
    'A live current-versus-baseline comparison and per-input change indicators.',
    'Four presets and reset to the original default scenario.',
    'Responsive dark interface with visible keyboard focus indicators and no required backend or API key.',
]),
('7. AI tools and development method', [
    'The initial project handoff states that the application had previously been developed using OpenCode. The original OpenCode conversation and prompt logs were not available in the files inspected for this report, so individual OpenCode prompts and contributions cannot be documented. The supplied Codex prompts and the project handover record show the later transition to Codex for diagnosis, repair, final verification, and report preparation.',
    'Codex inspected the existing project rather than replacing it. The first recorded work repaired broken application state and rendering, made presets and reset functional, corrected graph dependencies, and produced a running preview. A later development pass added input validation, handled the zero-budget gauge case, clarified the illustrative nature of the model, removed the remote font request, and added six built-in Node tests. The final pass in this conversation is focused on preparing this report and prompt book; it did not modify the working application.',
    'Prompt engineering evidence is provided separately in the prompt book. It reproduces the three available user prompts from the attached text files and summarizes outcomes supported by the project and recorded checks.']),
('8. Testing methodology and results', [
    'The simulation tests use Node’s built-in node:test and strict assertions. They cover baseline totals, exact budget deficit, deterministic pressure and risk values, range extremes and zero budget, rejection of invalid values, and variable ranges and updates. The records below distinguish automated test output from browser interaction checks performed in earlier project tasks.',
]),
('9. Results', [
    'For the default scenario, the model calculates catering expenditure of ₹120,000, total expenditure of ₹355,000, and ₹145,000 remaining from the ₹500,000 budget. It calculates 50% preparation pressure, 45% schedule risk, and “good” readiness. These values were observed in the rendered application during prior browser verification.',
    'The browser record also shows the double-guests preset changing 150 guests to 300 and total expenditure to ₹475,000; the expensive-venue preset creating a ₹5,000 budget deficit; the prep-time preset changing 45 days to 15 days, 83% pressure, 62% risk, and “at-risk” readiness; the budget preset increasing budget to ₹600,000; and reset restoring the baseline. The later production-preview check confirmed initial default values and exercised double-guests and reset.']),
('10. Limitations and future improvements', [
    'The model’s rules are illustrative and intentionally simplified. It does not include vendor quotes, taxes, deposits, transportation routes, task dependencies, uncertainty, or probabilistic estimates. Its readiness categories are rule-based labels, not validated recommendations. The graph represents only dependencies encoded in the application.',
    'Browser console inspection and a complete graph keyboard-selection test were not established in the available verification record. No user research, accessibility certification, performance benchmark, or public deployment was performed. The report format PDF was not available; the supplied image only partially shows a report cover.',
    'Possible future work includes completing the university-format review, capturing genuine application screenshots, testing all controls and graph interactions across screen sizes, adding scenario save/export only if needed, and conducting a documented user evaluation if the course requires it.']),
('11. Deployment readiness', [
    'This is a frontend-only Vite application. The production build output contains an HTML entry point and hashed CSS and JavaScript assets under dist. The current app has no application-level localhost URL dependency, server routes, or API credentials. A production preview was opened locally and rendered the dashboard. Deployment is not complete and no public URL exists.',
    'A straightforward manual route is to run npm run build and upload the generated dist folder to Netlify Drop. Netlify’s official instructions describe drag-and-drop publishing for a built output folder; its documentation says a manual deploy can produce a public URL without an account, while an account makes the deployment easier to claim and manage. Follow the current instructions at https://docs.netlify.com/start/quickstarts/netlify-drop-quickstart/ and https://docs.netlify.com/deploy/create-deploys/. Check the generated page and preset interactions at the resulting URL.']),
('12. Conclusion', [
    'The Butterfly Effect is a functioning interactive demonstration of a deterministic event-planning model. It combines editable scenario inputs, calculated outputs, an explicit dependency graph, presets, reset, and a fixed-baseline comparison. Available records show successful build, lint, simulation tests, and browser checks. The main submission dependencies are the complete university format, student identity details, genuine screenshots if required, and final manual review.']),
('References', [
    'Project source: src/App.tsx, src/simulation.ts, src/index.css, tests/simulation.test.mjs, package.json, and package-lock.json. Accessed during preparation of this report.',
    'The Butterfly Effect project handover: PROJECT_HANDOVER.md. Used as a record of project scope and prior checks; checked against source files and the conversation’s observed test output where available.',
    'Netlify Docs, “Netlify Drop Quickstart” and “Create deploys.” https://docs.netlify.com/start/quickstarts/netlify-drop-quickstart/ and https://docs.netlify.com/deploy/create-deploys/. Accessed 29 October 2026.',
    'University report-format cover image supplied by the student: IPE_Project_report_Format.docx.png. The complete format PDF or DOCX was not supplied, so this image was insufficient to establish required sections or detailed formatting.',
]),
]

tests = [
('Baseline totals and budget reconcile', '₹120,000 catering; ₹355,000 total; ₹145,000 remaining; zero deficit', 'Assertions passed for all four expected values', 'PASS — automated, previous final development pass'),
('Budget deficit is exact overspend', '₹855,000 spend; ₹355,000 negative remaining; ₹355,000 deficit', 'Assertions passed for all expected values', 'PASS — automated, previous final development pass'),
('Pressure, risk, determinism', 'At 7 days: 92% pressure and 66% risk; repeated calculation matches', 'Assertions passed; readiness expectation corrected to at-risk after first test run', 'PASS — automated, final run'),
('Range extremes and zero budget', 'Finite outputs; ₹6,200,000 max-range spend; ₹6,150,000 deficit; zero budget handled', 'Assertions passed', 'PASS — automated, final run'),
('Invalid values', 'NaN, infinity, and negative scenario input rejected', 'Expected RangeError assertions passed', 'PASS — automated, final run'),
('Ranges and update helper', 'Guest range is 20–1,000 step 10; update preserves other fields', 'Assertions passed', 'PASS — automated, final run'),
('Browser presets and reset', 'Presets update the intended input and derived values; reset restores baseline', 'Guest, venue, prep-time, budget presets and reset observed across prior and later browser checks', 'PASS — browser, prior task records'),
('Production preview', 'Static production bundle serves the app and can interact', 'Rendered at localhost:4173; double-guests preset and reset observed', 'PASS — browser, prior final development pass'),
('Console errors and graph keyboard selection', 'No runtime errors; keyboard selection works', 'No browser-console interface available; complete node-focus interaction not established', 'NOT VERIFIED'),
]

screenshots = [
('Figure 1. Application dashboard', 'Show the actual app at its default scenario with the scenario controls and Plan Outlook results visible. Open http://127.0.0.1:5173/ or run the local command in the handover section; reset to baseline; capture a screenshot; insert it in Section 6 after the feature list.'),
('Figure 2. Ripple map with an active change', 'Apply Double the guests and show the 300 guest value, recalculated spend, and highlighted graph path. Capture the actual rendered app and insert it in Section 5.1 after the workflow description.'),
('Figure 3. Baseline comparison', 'Show the current scenario beside the unchanged baseline and a visible change summary. Capture the actual rendered app after applying one preset and insert it in Section 9 after the results paragraph.'),
]

personal_fields = [
('Student name', '[STUDENT ACTION REQUIRED: Enter your full name.]'),
('GR number / Roll number', '[STUDENT ACTION REQUIRED: Enter your GR number or roll number using the university’s required field name.]'),
('Division', '[STUDENT ACTION REQUIRED: Enter your division.]'),
('College / institute', '[STUDENT ACTION REQUIRED: Enter the official college or institute name.]'),
('University', '[STUDENT ACTION REQUIRED: Enter the official university name if required.]'),
('Faculty / guide', '[STUDENT ACTION REQUIRED: Enter the faculty or guide name if required by the official format.]'),
('Academic year / semester', '[STUDENT ACTION REQUIRED: Enter the academic year and semester if required.]'),
]

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); el = OxmlElement('w:tblHeader'); el.set(qn('w:val'), 'true'); trPr.append(el)

def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    txt = OxmlElement('w:t'); txt.text = '1'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instr, sep, txt, end])

def setup_doc(doc, title):
    sec = doc.sections[0]
    sec.top_margin = Inches(.72); sec.bottom_margin = Inches(.7); sec.left_margin = Inches(.8); sec.right_margin = Inches(.8)
    styles = doc.styles
    normal = styles['Normal']; normal.font.name = 'Aptos'; normal.font.size = Pt(10); normal.font.color.rgb = RGBColor(42, 51, 59)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.12
    for name, size, color in [('Title', 26, '15242D'), ('Heading 1', 17, '12695E'), ('Heading 2', 13, '266E66'), ('Heading 3', 11, '344A52')]:
        s = styles[name]; s.font.name = 'Aptos Display'; s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(11); s.paragraph_format.space_after = Pt(5); s.paragraph_format.keep_with_next = True
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('THE BUTTERFLY EFFECT  •  IPE PROJECT  |  Page '); add_page_field(footer)
    doc.core_properties.title = title
    doc.core_properties.subject = 'University IPE project documentation'

def doc_cover(doc, title, subtitle=''):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(72); p.paragraph_format.space_after = Pt(22)
    r = p.add_run('PROJECT REPORT'); r.bold = True; r.font.size = Pt(17); r.font.color.rgb = RGBColor(25, 43, 50)
    p = doc.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(8)
    p.add_run(title)
    if subtitle:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run(subtitle); rr.italic = True; rr.font.size = Pt(12); rr.font.color.rgb = RGBColor(75, 103, 105)
    doc.add_paragraph('')
    t = doc.add_table(rows=0, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    t.columns[0].width = Inches(2.0); t.columns[1].width = Inches(4.2)
    for label, value in personal_fields:
        cells = t.add_row().cells; cells[0].text = label; cells[1].text = value
        cells[0].width = Inches(2.0); cells[1].width = Inches(4.2)
        for c in cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                for run in para.runs: run.font.size = Pt(9)
        shade(cells[0], 'E8F1EF'); shade(cells[1], 'FFF2D9')
        for run in cells[0].paragraphs[0].runs: run.bold = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(25)
    r = p.add_run('[STUDENT ACTION REQUIRED: Check the complete university format and page order. The only supplied format item was a small cover image.]'); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(175, 42, 34)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(15)
    p.add_run('Prepared 29 October 2026').font.size = Pt(9)
    doc.add_page_break()

def add_doc_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers)); table.style = 'Table Grid'; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])
    for i, text in enumerate(headers):
        c = table.rows[0].cells[i]; c.text = text; shade(c, '164D4A')
        for run in c.paragraphs[0].runs: run.bold = True; run.font.color.rgb = RGBColor(255,255,255); run.font.size = Pt(8)
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text); cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2 == 0: shade(cells[i], 'F1F5F4')
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs: run.font.size = Pt(7.5)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths): row.cells[i].width = Inches(width)
    return table

def add_placeholder(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(.15); p.paragraph_format.right_indent = Inches(.15)
    p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(10)
    pPr = p._p.get_or_add_pPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), 'FCE5E3'); pPr.append(shd)
    r = p.add_run('[STUDENT ACTION REQUIRED: ' + text + ']'); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(157, 32, 29)

def build_report_docx():
    doc = Document(); setup_doc(doc, 'The Butterfly Effect IPE Project Report')
    doc_cover(doc, 'The Butterfly Effect', 'Interactive Event Planning Scenario Demonstrator')
    doc.add_heading('Submission format note', 1)
    doc.add_paragraph('The report-format PDF referred to in the project request was not present in D:\\butterfly-effect. The supplied file is a small cover image and does not make the required chapters, declarations, table of contents, page order, or detailed formatting legible. This report is a complete technical draft based on available project evidence, not a claim of exact university-format compliance.')
    add_placeholder(doc, 'Obtain the official full report-format PDF or editable template and verify the required front matter, section titles, page order, declarations, and formatting. Adjust this report to match it before submission.')
    doc.add_heading('Student details', 1)
    add_doc_table(doc, ['Required field', 'Value to be completed'], personal_fields)
    doc.add_heading('Contents', 1)
    for title, _ in report_sections: doc.add_paragraph(title, style='List Bullet')
    doc.add_paragraph('Testing results', style='List Bullet'); doc.add_paragraph('Screenshot evidence', style='List Bullet')
    doc.add_page_break()
    for title, paragraphs in report_sections:
        doc.add_heading(title, 1)
        for text in paragraphs:
            if title == '6. Implemented features': doc.add_paragraph(text, style='List Bullet')
            else: doc.add_paragraph(text)
        if title == '8. Testing methodology and results':
            add_doc_table(doc, ['Test', 'Expected outcome', 'Observed outcome', 'Status'], tests, [1.4, 1.75, 1.8, 1.3])
        if title == '6. Implemented features':
            for cap, instruction in screenshots[:1]:
                doc.add_heading('Figure 1. Application dashboard evidence', 2); add_placeholder(doc, instruction)
        if title == '5.1 Workflow':
            doc.add_heading('Figure 2. Ripple map evidence', 2); add_placeholder(doc, screenshots[1][1])
        if title == '9. Results':
            doc.add_heading('Figure 3. Baseline comparison evidence', 2); add_placeholder(doc, screenshots[2][1])
    doc.add_page_break(); doc.add_heading('Appendix A. Declarations and acknowledgements', 1)
    doc.add_paragraph('The supplied cover image does not establish whether the university requires a certificate, originality declaration, or acknowledgements page. Do not use invented institutional wording or signatures.')
    add_placeholder(doc, 'Insert the exact required certificate, declaration, and acknowledgement pages from the official university format. Complete and sign them personally where required.')
    doc.add_heading('Appendix B. Local project access', 1)
    doc.add_paragraph('Run the application from PowerShell with:')
    p = doc.add_paragraph(); r = p.add_run('cd D:\\butterfly-effect\nnpm run dev'); r.font.name = 'Consolas'; r.font.size = Pt(10)
    doc.add_paragraph('The application is available at http://127.0.0.1:5173/ while the development server is running.')
    path = OUT / 'Division_GR No_Name.docx'; doc.save(path); return path

def build_prompt_docx():
    doc = Document(); setup_doc(doc, 'The Butterfly Effect IPE Prompt Book')
    doc_cover(doc, 'The Butterfly Effect', 'Prompt Book and Development Record')
    doc.add_heading('Record scope', 1)
    doc.add_paragraph('This prompt book reproduces the three user prompts available in the supplied Codex attachment files. It distinguishes those texts from summaries of outcomes. The initial project handoff says development previously used OpenCode, but no OpenCode transcript or original prompt log was available; those prompts are not reconstructed or invented.')
    doc.add_heading('Development record', 1)
    add_doc_table(doc, ['No.', 'Purpose and tool', 'Outcome supported by available record'], [
        ('1', 'Diagnose, repair, verify app — Codex; handoff notes an earlier OpenCode phase.', 'Codex inspected and repaired the existing Vite/React app, made interactions functional, and recorded browser verification. Original OpenCode prompts unavailable.'),
        ('2', 'Final application pass — Codex.', 'Input validation, zero-budget gauge, illustrative-model wording, removal of remote font dependency, and six Node tests. Build, lint, tests, and production preview were verified in the conversation.'),
        ('3', 'Report and prompt-book preparation — Codex.', 'Current report and this prompt book prepared. Full university format remains unavailable; identity fields and screenshot evidence remain student actions.'),
    ])
    for i, (heading, path, purpose) in enumerate(PROMPTS, 1):
        doc.add_page_break(); doc.add_heading(f'{heading}', 1)
        doc.add_paragraph(f'AI tool: Codex desktop assistant. The project handoff reports an earlier OpenCode phase, but this exact prompt was supplied for Codex work. Purpose: {purpose}')
        doc.add_paragraph('Intended outcome: follow the request below while preserving the existing application and reporting evidence accurately.')
        if path.exists():
            prompt_text = path.read_text(encoding='utf-8').strip()
            for block in re.split(r'\n\s*\n', prompt_text):
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5)
                r = p.add_run(block); r.font.size = Pt(8.5); r.font.name = 'Aptos'
        else:
            add_placeholder(doc, f'The source prompt file {path.name} is not available; student must supply it if a verbatim prompt record is required.')
        if i == 1:
            doc.add_heading('Recorded outcome and verification', 2)
            doc.add_paragraph('The application was inspected and repaired rather than replaced. The later project record shows successful build and lint, browser testing of all presets, reset, and slider interaction, and a visible local preview. The detailed repair history predates this report task and is summarized from the project handover and current conversation record.')
        elif i == 2:
            doc.add_heading('Recorded outcome and verification', 2)
            doc.add_paragraph('The final development pass added finite/non-negative validation, zero-budget gauge behavior, an explicit illustrative-model description, local font fallbacks, and six Node-native tests. The observed final checks were: production build exit 0; Oxlint exit 0 with no diagnostics; six tests passed and zero failed; production preview at 127.0.0.1:4173 rendered and its double-guests preset and reset were exercised.')
            doc.add_paragraph('Limitations in the record: no browser console inspection, no full graph keyboard node-selection test, and no public deployment.')
        else:
            doc.add_heading('Recorded outcome and verification', 2)
            doc.add_paragraph('The current document task inspected project files and prior verification records. The project directory did not contain the requested university-format PDF; the supplied format image shows only a cover cue and does not establish full report requirements. The report and prompt book therefore contain visible student-action placeholders. No public publishing or fabricated evidence was performed.')
    doc.add_page_break(); doc.add_heading('Prompt history that remains unavailable', 1)
    doc.add_paragraph('The source project handoff identifies OpenCode as an earlier development tool, but no original OpenCode prompt text, conversation transcript, or tool logs were found in the inspected project files or supplied attachments. The prompts above are the available user instructions for the three Codex work phases. The student should provide any original OpenCode records if the university requires them.')
    path = OUT / 'Prompt_Book_The_Butterfly_Effect.docx'; doc.save(path); return path

# PDF from the same source data: native text and tables, not rasterized screenshots.
def build_pdf(docx_path, is_prompt=False):
    pdf_path = docx_path.with_suffix('.pdf')
    if FONT and Path(FONT).exists():
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        pdfmetrics.registerFont(TTFont('Body', FONT)); pdfmetrics.registerFont(TTFont('Bold', FONT))
        font_name = 'Body'
    else: font_name = 'Helvetica'
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('DocTitle', parent=styles['Title'], fontName=font_name, fontSize=24, leading=30, textColor=colors.HexColor('#15242D'), alignment=TA_CENTER, spaceAfter=10)
    h1 = ParagraphStyle('H1x', parent=styles['Heading1'], fontName=font_name, fontSize=15, leading=19, textColor=colors.HexColor('#12695E'), spaceBefore=11, spaceAfter=6, keepWithNext=True)
    h2 = ParagraphStyle('H2x', parent=styles['Heading2'], fontName=font_name, fontSize=11, leading=14, textColor=colors.HexColor('#266E66'), spaceBefore=7, spaceAfter=4, keepWithNext=True)
    body = ParagraphStyle('Bodyx', parent=styles['BodyText'], fontName=font_name, fontSize=9, leading=13, textColor=colors.HexColor('#2A333B'), spaceAfter=6)
    small = ParagraphStyle('Smallx', parent=body, fontSize=7.3, leading=9.3, spaceAfter=1)
    callout = ParagraphStyle('Callout', parent=body, fontName=font_name, fontSize=8.5, leading=12, textColor=colors.HexColor('#9D201D'), backColor=colors.HexColor('#FCE5E3'), borderPadding=6, spaceBefore=5, spaceAfter=9)
    story=[]
    esc=lambda s: str(s).replace('₹', 'INR ').replace('&','&amp;').replace('<','&lt;').replace('>','&gt;').replace('\n','<br/>')
    def table_pdf(headers, rows, widths, fs=7):
        header_style=ParagraphStyle('TableHeader', parent=small, textColor=colors.HexColor('#153F3B'), fontName=font_name, fontSize=7.3, leading=9.3)
        data=[[Paragraph(f'<b>{esc(x)}</b>', header_style) for x in headers]]
        data += [[Paragraph(esc(x), small) for x in row] for row in rows]
        t=Table(data, colWidths=[w*inch for w in widths], repeatRows=1, hAlign='LEFT')
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DCEBE8')),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#BAC8C7')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F1F5F4')]),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]))
        return t
    def cover(title, subtitle):
        story.extend([Spacer(1,45*mm),Paragraph('PROJECT REPORT', ParagraphStyle('coverlabel',parent=body,fontName=font_name,fontSize=16,alignment=TA_CENTER,spaceAfter=17)),Paragraph(esc(title),title_style),Paragraph(esc(subtitle),ParagraphStyle('coversub',parent=body,fontName=font_name,fontSize=12,alignment=TA_CENTER,textColor=colors.HexColor('#4B6769'),spaceAfter=22))])
        story.append(table_pdf(['Required field','Value to be completed'],personal_fields,[2.0,4.6],8))
        story.extend([Spacer(1,14),Paragraph('STUDENT ACTION REQUIRED: Check the complete university format and page order. The supplied item was a small cover image only.',callout),Paragraph('Prepared 29 October 2026',ParagraphStyle('coverdate',parent=body,alignment=TA_CENTER)) ,PageBreak()])
    if is_prompt:
        cover('The Butterfly Effect','Prompt Book and Development Record')
        story.append(Paragraph('Record scope',h1)); story.append(Paragraph('This prompt book reproduces the three user prompts available in the supplied Codex attachment files. The initial handoff says the project previously used OpenCode, but no original OpenCode transcript or prompts were available. Those prompts are not reconstructed.',body))
        story.append(Paragraph('Development record',h1)); story.append(table_pdf(['No.','Purpose and tool','Recorded outcome'],[
            ('1','Diagnosis and repair — Codex; earlier OpenCode reported.','Existing app repaired and browser verified. OpenCode prompt records unavailable.'),
            ('2','Final application pass — Codex.','Validation, zero-budget gauge, illustrative wording, local fonts, six tests; build/lint/tests/production preview verified.'),
            ('3','Report and prompt book — Codex.','Documents prepared; full format PDF, personal information, and screenshots remain outstanding.')],[.45,2.2,3.95],7))
        for i,(heading,path,purpose) in enumerate(PROMPTS,1):
            story.append(PageBreak()); story.append(Paragraph(esc(heading),h1)); story.append(Paragraph('AI tool: Codex desktop assistant. Purpose: '+esc(purpose),body)); story.append(Paragraph('Actual user prompt text',h2))
            text=path.read_text(encoding='utf-8').strip() if path.exists() else f'[PROMPT UNAVAILABLE: {path.name}]'
            for block in re.split(r'\n\s*\n',text): story.append(Paragraph(esc(block),small)); story.append(Spacer(1,3))
            story.append(Paragraph('Recorded outcome and verification',h2))
            outcome=[
                'The existing application was inspected and repaired. Available records show production build and lint success, browser checks for presets/reset/sliders, and a local preview. This prompt predates the report task.',
                'This pass added input validation, zero-budget gauge behavior, an illustrative-model description, local font fallback, and six tests. Observed checks: build exit 0; Oxlint exit 0; six tests passed and zero failed; production preview rendered, and double-guests/reset were exercised. Console inspection, full graph keyboard testing, and public deployment remain incomplete.',
                'The project folder did not contain the requested university PDF. The supplied cover image did not establish the full report structure. Documents were prepared with visible student-action placeholders; no external deployment was made.',
            ][i-1]
            story.append(Paragraph(esc(outcome),body))
        story.extend([PageBreak(),Paragraph('Prompt history that remains unavailable',h1),Paragraph('The project handoff identifies OpenCode as an earlier development tool, but no OpenCode prompt text or transcripts were found in the inspected project or supplied attachments. Provide any original OpenCode records if the university requires them.',body)])
    else:
        cover('The Butterfly Effect','Interactive Event Planning Scenario Demonstrator')
        story.append(Paragraph('Submission format note',h1)); story.append(Paragraph('The report-format PDF referenced in the request was not present in D:\\butterfly-effect. The supplied file is a small cover image and does not make chapter requirements, declarations, page order, or detailed formatting legible. This is a complete technical draft, not a claim of exact university-format compliance.',body)); story.append(Paragraph('STUDENT ACTION REQUIRED: Obtain the full official format and adapt the report before submission.',callout))
        story.append(Paragraph('Student details',h1)); story.append(table_pdf(['Required field','Value to be completed'],personal_fields,[2.0,4.6],8))
        story.append(Paragraph('Contents',h1)); story.append(Paragraph(' | '.join(esc(x[0]) for x in report_sections)+' | Testing results | Screenshot evidence',body)); story.append(PageBreak())
        for title, paragraphs in report_sections:
            story.append(Paragraph(esc(title),h1))
            for text in paragraphs:
                if title=='6. Implemented features': story.append(Paragraph('• '+esc(text),body))
                else: story.append(Paragraph(esc(text),body))
            if title=='8. Testing methodology and results': story.append(table_pdf(['Test','Expected','Observed','Status'],tests,[1.3,1.75,2.0,1.45],6.6))
            if title in ('5.1 Workflow','6. Implemented features','9. Results'):
                cap,instr=screenshots[{'5.1 Workflow':1,'6. Implemented features':0,'9. Results':2}[title]]
                story.append(Paragraph(esc(cap),h2)); story.append(Paragraph('STUDENT ACTION REQUIRED: '+esc(instr),callout))
        story.extend([PageBreak(),Paragraph('Appendix A. Declarations and acknowledgements',h1),Paragraph('The supplied cover image does not establish whether the university requires a certificate, originality declaration, or acknowledgements page. Do not use invented institutional wording or signatures.',body),Paragraph('STUDENT ACTION REQUIRED: Insert the exact required pages and personally complete/sign fields where required.',callout),Paragraph('Appendix B. Local project access',h1),Paragraph('From PowerShell: <font name="Courier">cd D:\\butterfly-effect</font>, then <font name="Courier">npm run dev</font>. Open http://127.0.0.1:5173/ while the dev server is running.',body)])
    def footer(canvas, doc):
        canvas.saveState(); canvas.setFont('Helvetica',7); canvas.setFillColor(colors.HexColor('#69777B')); canvas.drawCentredString(A4[0]/2,10*mm,f'THE BUTTERFLY EFFECT  |  IPE PROJECT  •  {doc.page}'); canvas.restoreState()
    pdf=SimpleDocTemplate(str(pdf_path),pagesize=A4,rightMargin=19*mm,leftMargin=19*mm,topMargin=18*mm,bottomMargin=18*mm,title='The Butterfly Effect IPE Documentation',author='')
    pdf.build(story,onFirstPage=footer,onLaterPages=footer)
    return pdf_path

if __name__ == '__main__':
    report=build_report_docx(); prompt=build_prompt_docx()
    report_pdf=build_pdf(report,False); prompt_pdf=build_pdf(prompt,True)
    print('\n'.join(map(str,[report,prompt,report_pdf,prompt_pdf])))
